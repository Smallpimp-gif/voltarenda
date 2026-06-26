#!/usr/bin/env python3
# Импорт активных арендаторов из Excel-таблицы в data/bot/tenants.json.
#
# Таблица «Платежи арендаторов» — источник истины (её ведёт оператор).
# Бот же читает только JSON. Этот скрипт — мост: запусти после правок
# в Excel, и список в боте обновится.
#
# Использование:
#   python3 bot/import-tenants.py "/путь/к/Платежи_арендаторов.xlsx"
#   (без аргумента берёт ~/Downloads/Платежи_арендаторов.xlsx)
#
# Берём только строки со статусом «активен» — закрытым договорам
# напоминания не нужны. Поля username (Telegram) НЕ трогаем: их
# заполняет оператор через бота, перезапись стёрла бы привязки.

import json
import os
import sys
import unicodedata

try:
    import openpyxl
except ImportError:
    sys.exit("Нужен openpyxl:  pip3 install openpyxl")

# Транслитерация фамилии в стабильный латинский id (для callback-данных
# Telegram — там лимит 64 байта, кириллица ест его быстро).
_TRANSLIT = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e",
    "ж": "zh", "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m",
    "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "у": "u",
    "ф": "f", "х": "h", "ц": "ts", "ч": "ch", "ш": "sh", "щ": "sch",
    "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya",
}


def slugify(name: str) -> str:
    # Отрезаем уточнение в скобках: «Соколов (U7 выкуп)» → «Соколов».
    base = name.split("(")[0].strip().lower()
    out = []
    for ch in base:
        if ch in _TRANSLIT:
            out.append(_TRANSLIT[ch])
        elif ch.isalnum():
            out.append(unicodedata.normalize("NFKD", ch).encode("ascii", "ignore").decode())
        else:
            out.append("-")
    slug = "".join(out).strip("-")
    while "--" in slug:
        slug = slug.replace("--", "-")
    return slug or "tenant"


def main() -> None:
    default = os.path.expanduser("~/Downloads/Платежи_арендаторов.xlsx")
    path = sys.argv[1] if len(sys.argv) > 1 else default
    if not os.path.exists(path):
        sys.exit(f"Файл не найден: {path}")

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Платежи"] if "Платежи" in wb.sheetnames else wb.worksheets[0]

    rows = list(ws.iter_rows(values_only=True))
    # Находим строку заголовков (где первая ячейка == «Арендатор»).
    header_idx = next(i for i, r in enumerate(rows) if r and r[0] == "Арендатор")
    cols = {name: idx for idx, name in enumerate(rows[header_idx]) if name}

    def cell(row, col_name):
        idx = cols.get(col_name)
        return row[idx] if idx is not None and idx < len(row) else None

    tenants = []
    seen = {}
    for row in rows[header_idx + 1:]:
        name = row[0]
        if not name or name == "ИТОГО":
            continue
        status = cell(row, "Статус")
        if status != "активен":
            continue

        start = cell(row, "Дата старта")
        start_iso = start.strftime("%Y-%m-%d") if hasattr(start, "strftime") else str(start)

        ttype = cell(row, "Тип")  # «аренда» | «выкуп»
        weekly = cell(row, "Платёж/нед")
        buyout_weeks = cell(row, "Срок выкупа")
        buyout_weeks = int(buyout_weeks) if isinstance(buyout_weeks, (int, float)) else None

        base_id = slugify(name)
        uid = base_id
        if base_id in seen:
            seen[base_id] += 1
            uid = f"{base_id}-{seen[base_id]}"
        else:
            seen[base_id] = 1

        tenants.append({
            "id": uid,
            "name": name,
            "contract": cell(row, "Договор") or "",
            "type": ttype,
            "weekly": int(weekly) if isinstance(weekly, (int, float)) else 0,
            "startDate": start_iso,
            "buyoutWeeks": buyout_weeks,
            # Telegram аккаунт арендатора. Заполняет оператор (через бота
            # или вручную). Пусто → привязка вручную после нажатия «Старт».
            "telegramUsername": "",
        })

    out_dir = os.path.join(os.path.dirname(__file__), "..", "data", "bot")
    out_dir = os.path.abspath(out_dir)
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "tenants.json")

    # Сохраняем уже введённые оператором username при переимпорте.
    if os.path.exists(out_path):
        with open(out_path, encoding="utf-8") as f:
            old = {t["id"]: t for t in json.load(f)}
        for t in tenants:
            if t["id"] in old and old[t["id"]].get("telegramUsername"):
                t["telegramUsername"] = old[t["id"]]["telegramUsername"]

    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(tenants, f, ensure_ascii=False, indent=2)

    print(f"OK: {len(tenants)} активных арендаторов → {out_path}")
    for t in tenants:
        wk = f"{t['buyoutWeeks']} нед" if t["buyoutWeeks"] else "бессрочно"
        print(f"  {t['id']:<16} {t['name']:<24} {t['type']:<7} {t['weekly']}₽/нед  с {t['startDate']}  ({wk})")


if __name__ == "__main__":
    main()
